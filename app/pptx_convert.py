"""PDF ↔ PowerPoint and related Office format helpers."""

from __future__ import annotations

import io
import logging
import os
import shutil
from pathlib import Path
from typing import Callable

from fastapi import HTTPException

logger = logging.getLogger("pdfsuite")

# LibreOffice filter names for reliable cross-format conversion
LO_FILTERS = {
    "pdf": "pdf:writer_pdf_Export",
    "docx": "docx:MS Word 2007 XML",
    "pptx": "pptx:Impress MS PowerPoint 2007 XML",
    "xlsx": "xlsx:Calc MS Excel 2007 XML",
    "odt": "odt:writer8",
    "odp": "odp:impress8",
    "ods": "ods:calc8",
}

# Allowed source → target pairs (excluding PDF→Office which uses dedicated engines).
# Word ↔ Excel is via LibreOffice (best when the Word file has tables).
OFFICE_FORMAT_MAP: dict[str, set[str]] = {
    ".ppt": {"pdf", "pptx", "odp"},
    ".pptx": {"pdf", "docx", "xlsx", "odp"},
    ".doc": {"pdf", "docx", "xlsx", "odt", "pptx"},
    ".docx": {"pdf", "pptx", "xlsx", "odt"},
    ".xls": {"pdf", "xlsx", "ods", "docx"},
    ".xlsx": {"pdf", "pptx", "ods", "docx"},
    ".odp": {"pdf", "pptx", "docx"},
    ".odt": {"pdf", "docx", "pptx", "xlsx"},
    ".ods": {"pdf", "xlsx", "docx"},
}

SOURCE_LABELS: dict[str, str] = {
    ".pdf": "PDF",
    ".doc": "Word (legacy)",
    ".docx": "Word",
    ".ppt": "PowerPoint (legacy)",
    ".pptx": "PowerPoint",
    ".xls": "Excel (legacy)",
    ".xlsx": "Excel",
    ".odt": "OpenDocument Text",
    ".odp": "OpenDocument Presentation",
    ".ods": "OpenDocument Spreadsheet",
}

TARGET_LABELS: dict[str, str] = {
    "pdf": "PDF",
    "docx": "Word (.docx)",
    "pptx": "PowerPoint (.pptx)",
    "xlsx": "Excel (.xlsx)",
    "odt": "OpenDocument Text (.odt)",
    "odp": "OpenDocument Presentation (.odp)",
    "ods": "OpenDocument Spreadsheet (.ods)",
    "word": "Word (.docx)",
    "excel": "Excel (.xlsx) — extract tables",
    "word_layout": "Word layout mode",
}

# Extra hints for trickier conversions
TARGET_HINTS: dict[tuple[str, str], str] = {
    (".docx", "xlsx"): "Best results when Word contains tables",
    (".doc", "xlsx"): "Best results when Word contains tables",
    (".odt", "xlsx"): "Best results when the document contains tables",
    (".xlsx", "docx"): "Spreadsheet exported as a Word document",
    (".xls", "docx"): "Spreadsheet exported as a Word document",
    (".pdf", "excel"): "Extract tables from the PDF into a spreadsheet",
    (".pdf", "pptx"): "One image slide per PDF page",
    (".pdf", "word"): "Editable text document",
}


def normalize_ext(name_or_ext: str) -> str:
    raw = (name_or_ext or "").strip().lower()
    if not raw:
        return ""
    if "/" in raw or "\\" in raw:
        raw = Path(raw).suffix.lower()
    if not raw.startswith(".") and "." in raw:
        raw = Path(raw).suffix.lower()
    if raw and not raw.startswith("."):
        raw = f".{raw}"
    return raw


def detect_source(filename: str) -> dict:
    """
    Detect source type from filename and list valid conversion targets.
    Fool-proof routing for the Smart Convert UI.
    """
    ext = normalize_ext(filename)
    if not ext:
        return {
            "ok": False,
            "error": "Could not detect file type. Use a named file with an extension.",
            "ext": "",
            "label": "Unknown",
            "targets": [],
        }

    if ext == ".pdf":
        targets = [
            {
                "id": "word",
                "label": "Word (.docx)",
                "route": "pdf_word",
                "hint": TARGET_HINTS.get((".pdf", "word"), "Editable text document"),
            },
            {
                "id": "excel",
                "label": "Excel (.xlsx)",
                "route": "pdf_excel",
                "hint": TARGET_HINTS.get((".pdf", "excel"), "Extract tables into a spreadsheet"),
            },
            {
                "id": "pptx",
                "label": "PowerPoint (.pptx)",
                "route": "pdf_pptx",
                "hint": TARGET_HINTS.get((".pdf", "pptx"), "One image slide per PDF page"),
            },
        ]
        return {
            "ok": True,
            "ext": ext,
            "label": SOURCE_LABELS[ext],
            "family": "pdf",
            "targets": targets,
            "supports_pages": True,
            "supports_word_mode": True,
        }

    if ext in OFFICE_FORMAT_MAP:
        targets = []
        for t in sorted(OFFICE_FORMAT_MAP[ext]):
            default_hint = f"{SOURCE_LABELS.get(ext, ext)} → {TARGET_LABELS.get(t, t)}"
            targets.append(
                {
                    "id": t,
                    "label": TARGET_LABELS.get(t, t.upper()),
                    "route": "office",
                    "hint": TARGET_HINTS.get((ext, t), default_hint),
                }
            )
        return {
            "ok": True,
            "ext": ext,
            "label": SOURCE_LABELS.get(ext, ext),
            "family": "office",
            "targets": targets,
            "supports_pages": False,
            "supports_word_mode": False,
        }

    return {
        "ok": False,
        "error": (
            f"Unsupported type '{ext}'. "
            "Use PDF, Word, Excel, PowerPoint, or OpenDocument files."
        ),
        "ext": ext,
        "label": "Unsupported",
        "targets": [],
    }


def pdf_to_pptx(
    input_path: str | Path,
    output_path: str | Path,
    indices: list[int],
    *,
    dpi: int = 144,
    on_progress: Callable[[int, int], None] | None = None,
    should_cancel: Callable[[], bool] | None = None,
) -> int:
    """
    Build a PowerPoint deck from PDF pages (one full-bleed image slide per page).
    Preserves visual layout better than text-only for slides.
    """
    try:
        import fitz
    except ImportError as exc:  # pragma: no cover
        raise HTTPException(status_code=503, detail="PyMuPDF is required for PDF → PowerPoint.") from exc

    try:
        from pptx import Presentation
        from pptx.util import Emu, Inches
    except ImportError as exc:  # pragma: no cover
        raise HTTPException(status_code=503, detail="python-pptx is required for PowerPoint export.") from exc

    src = fitz.open(str(input_path))
    try:
        if not indices:
            raise HTTPException(status_code=400, detail="No pages selected.")

        # Use first page size for slide aspect (points → inches)
        first = src[indices[0]] if indices[0] < src.page_count else src[0]
        page_w_pt = float(first.rect.width)
        page_h_pt = float(first.rect.height)
        # 72 pt = 1 inch
        slide_w_in = page_w_pt / 72.0
        slide_h_in = page_h_pt / 72.0
        # Clamp ridiculous sizes
        slide_w_in = max(5.0, min(20.0, slide_w_in))
        slide_h_in = max(4.0, min(20.0, slide_h_in))

        prs = Presentation()
        prs.slide_width = Inches(slide_w_in)
        prs.slide_height = Inches(slide_h_in)
        blank = prs.slide_layouts[6]  # blank

        zoom = dpi / 72.0
        mat = fitz.Matrix(zoom, zoom)
        total = len(indices)
        written = 0

        for n, idx in enumerate(indices, start=1):
            if should_cancel and should_cancel():
                raise RuntimeError("cancelled")
            if idx < 0 or idx >= src.page_count:
                continue
            page = src[idx]
            pix = page.get_pixmap(matrix=mat, alpha=False)
            img_bytes = pix.tobytes("jpeg")

            slide = prs.slides.add_slide(blank)
            stream = io.BytesIO(img_bytes)
            # Full-bleed image
            slide.shapes.add_picture(
                stream,
                Emu(0),
                Emu(0),
                width=prs.slide_width,
                height=prs.slide_height,
            )
            written += 1
            if on_progress and (n == 1 or n == total or n % 3 == 0):
                on_progress(n, total)

        if written == 0:
            raise HTTPException(status_code=400, detail="No pages converted to PowerPoint.")
        if on_progress:
            on_progress(total, total)
        prs.save(str(output_path))
        return written
    finally:
        src.close()


def libreoffice_convert(
    *,
    libreoffice_path: str,
    input_path: str | Path,
    output_dir: str | Path,
    target_ext: str,
    run_subprocess,
    settings,
) -> str:
    """
    Convert an Office document to another format via LibreOffice.
    Returns path to the produced file.
    target_ext: 'pdf', 'docx', 'pptx', 'xlsx' (no dot).
    """
    target = target_ext.lower().lstrip(".")
    src = Path(input_path)
    out_dir = Path(output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    # LibreOffice convert-to can use filter or simple extension
    convert_arg = target
    if target in LO_FILTERS:
        # Simple extension is more portable across LO versions
        convert_arg = target

    run_subprocess(
        [
            libreoffice_path,
            "--headless",
            "--nologo",
            "--nofirststartwizard",
            "--convert-to",
            convert_arg,
            "--outdir",
            str(out_dir),
            str(src),
        ],
        settings,
        label=f"LibreOffice → {target}",
    )

    expected = out_dir / f"{src.stem}.{target}"
    if expected.is_file():
        return str(expected)

    # LO sometimes mangles names
    candidates = list(out_dir.glob(f"*.{target}"))
    if candidates:
        return str(candidates[0])

    raise HTTPException(
        status_code=500,
        detail=f"LibreOffice did not produce a .{target} file.",
    )


def media_type_for_ext(ext: str) -> str:
    e = ext.lower().lstrip(".")
    return {
        "pdf": "application/pdf",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "odt": "application/vnd.oasis.opendocument.text",
        "odp": "application/vnd.oasis.opendocument.presentation",
        "ods": "application/vnd.oasis.opendocument.spreadsheet",
    }.get(e, "application/octet-stream")
